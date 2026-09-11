import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Set

# Ensure parent directory (AI-Backend) and repo root are on sys.path for clean imports
BASE_DIR = Path(__file__).resolve().parent.parent
REPO_DIR = BASE_DIR.parent
for p in [str(BASE_DIR), str(REPO_DIR)]:
    if p not in sys.path:
        sys.path.insert(0, p)

from fastapi import APIRouter, HTTPException, Request, Response
from fastapi.routing import APIRoute

try:
    from fast_api.schemas import (
        AgentStateInfo,
        ChatRequest,
        ChatResponse,
        CreatedFileInfo,
    )
except (ImportError, ModuleNotFoundError):
    from schemas import (
        AgentStateInfo,
        ChatRequest,
        ChatResponse,
        CreatedFileInfo,
    )

from orchestration.orchestrator import build_graph


def sanitize_json_string(text: str) -> str:
    """
    Sanitizes unescaped Windows paths and invalid escape sequences in JSON payloads.
    Converts Windows paths like "C:\code\SIH_26_Repo\sample_files\students.csv"
    so standard JSON decoders can parse them without 'Invalid \escape' errors.
    """
    def fix_path(m):
        return m.group(0).replace('\\', '/')

    # Replace backslashes in Windows drive paths: "C:\..." or "D:\..."
    text = re.sub(r'"[a-zA-Z]:[\\/][^"\r\n]*"', fix_path, text)
    # Fix any remaining unescaped backslashes in strings
    text = re.sub(r'\\(?!["\\/bfnrt]|u[0-9a-fA-F]{4})', r'\\\\', text)
    return text


class WindowsPathJsonRoute(APIRoute):
    """
    Custom APIRoute that intercepts incoming JSON requests and auto-sanitizes
    unescaped Windows backslashes in file paths before Pydantic parsing.
    """
    def get_route_handler(self):
        original_route_handler = super().get_route_handler()

        async def custom_route_handler(request: Request) -> Response:
            req_content_type = request.headers.get("content-type", "")
            if "application/json" in req_content_type:
                raw_body = await request.body()
                if raw_body:
                    sanitized_text = sanitize_json_string(raw_body.decode("utf-8", errors="replace"))
                    sanitized_bytes = sanitized_text.encode("utf-8")

                    async def custom_receive():
                        return {"type": "http.request", "body": sanitized_bytes, "more_body": False}

                    request._receive = custom_receive
                    request._body = sanitized_bytes
            return await original_route_handler(request)

        return custom_route_handler


router = APIRouter(prefix="/api", tags=["Sovereign AI Workbench"], route_class=WindowsPathJsonRoute)

# Directories for uploads and generated outputs
UPLOAD_DIR = BASE_DIR / "uploads"
GENERATED_DIR = BASE_DIR / "generated_files"
OUTPUTS_DIR = BASE_DIR / "outputs"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
GENERATED_DIR.mkdir(parents=True, exist_ok=True)
OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)

# Supported document formats
SUPPORTED_DOC_EXTENSIONS = {".txt", ".pdf", ".docx", ".pptx", ".xlsx"}
GENERATED_DOC_EXTENSIONS = {".docx", ".xlsx", ".pptx", ".pdf", ".csv", ".png", ".jpg", ".txt"}

# Compile the LangGraph agent graph once at startup
graph = build_graph()


def _get_directory_files_snapshot(directories: List[Path]) -> Dict[Path, float]:
    """
    Returns a mapping of resolved Path -> modification time (mtime)
    for all files in the given directories.
    """
    snapshot: Dict[Path, float] = {}
    for d in directories:
        if d.exists():
            for f in d.glob("*"):
                if f.is_file():
                    try:
                        snapshot[f.resolve()] = f.stat().st_mtime
                    except Exception:
                        snapshot[f.resolve()] = 0.0
    return snapshot


def _get_directory_files(directory: Path) -> Set[Path]:
    """Returns a set of all existing file paths in a directory."""
    if not directory.exists():
        return set()
    return {f.resolve() for f in directory.glob("*") if f.is_file()}


def _extract_tools_used(messages: list) -> List[str]:
    """
    Extracts the unique names of tools that were executed during the agent's turn.
    Inspects ToolMessage instances and AIMessage tool_calls.
    """
    tools_used: List[str] = []

    for msg in messages:
        # Check ToolMessage name attribute
        name = getattr(msg, "name", None)
        if name and name not in tools_used:
            tools_used.append(name)

        # Check AIMessage tool_calls
        tool_calls = getattr(msg, "tool_calls", None)
        if tool_calls and isinstance(tool_calls, list):
            for tc in tool_calls:
                tc_name = tc.get("name") if isinstance(tc, dict) else getattr(tc, "name", None)
                if tc_name and tc_name not in tools_used:
                    tools_used.append(tc_name)

    return tools_used


def _ensure_in_output_directory(source_path: Path) -> Path:
    """
    Ensures that a generated file resides inside OUTPUTS_DIR.
    If it was created in the root or cwd, moves it to OUTPUTS_DIR.
    """
    resolved_source = source_path.resolve()
    target_path = OUTPUTS_DIR / resolved_source.name

    # If already in OUTPUTS_DIR or GENERATED_DIR, keep as is
    if resolved_source.parent in [OUTPUTS_DIR.resolve(), GENERATED_DIR.resolve()]:
        return resolved_source

    try:
        shutil.move(str(resolved_source), str(target_path))
        return target_path.resolve()
    except Exception:
        # Fallback to copy or return original if move fails
        try:
            shutil.copy2(str(resolved_source), str(target_path))
            return target_path.resolve()
        except Exception:
            return resolved_source


def _resolve_created_files(
    created_file_paths: List[str],
    before_files: Any,
    response_text: str,
    turn_start_time: Optional[float] = None
) -> List[CreatedFileInfo]:
    """
    Resolves created files by combining:
    1. Direct paths in AgentState['created_files'] (from DOCUMENT_CREATED tool outputs)
    2. Newly created or modified files in OUTPUTS_DIR, GENERATED_DIR, BASE_DIR, REPO_DIR, or CWD
    3. Document paths found in text responses
    Ensures all created files have their directory / file_directory / output_directory populated.
    """
    detected_files: List[CreatedFileInfo] = []
    seen_paths: Set[str] = set()

    # Convert before_files to dict of path -> mtime if it was a set or list
    before_snapshot: Dict[Path, float] = {}
    if isinstance(before_files, dict):
        before_snapshot = before_files
    elif isinstance(before_files, (set, list)):
        for item in before_files:
            p = Path(item).resolve() if not isinstance(item, Path) else item.resolve()
            before_snapshot[p] = p.stat().st_mtime if p.exists() else 0.0

    def _add_file(path_obj: Path):
        try:
            # Relocate to OUTPUTS_DIR if generated outside of output directories
            final_path = _ensure_in_output_directory(path_obj)
            resolved_str = str(final_path.resolve())
            dir_str = str(final_path.parent.resolve())

            if resolved_str not in seen_paths and final_path.is_file():
                seen_paths.add(resolved_str)
                size = final_path.stat().st_size if final_path.exists() else None
                detected_files.append(
                    CreatedFileInfo(
                        filename=final_path.name,
                        directory=dir_str,
                        file_directory=dir_str,
                        file_path=resolved_str,
                        output_directory=dir_str,
                        output_file_directory=dir_str,
                        download_url=f"/api/download/{final_path.name}",
                        size_bytes=size
                    )
                )
        except Exception:
            pass

    # 1. Direct paths reported in AgentState['created_files']
    for raw_path in (created_file_paths or []):
        if not raw_path:
            continue
        clean_path = str(raw_path).strip().strip("'\"")
        if clean_path.upper().startswith("DOCUMENT_CREATED:"):
            clean_path = clean_path.split(":", 1)[1].strip()
        candidate = Path(clean_path)

        # Search candidates in common locations if not found directly
        candidates_to_try = [
            candidate,
            OUTPUTS_DIR / candidate.name,
            GENERATED_DIR / candidate.name,
            BASE_DIR / candidate,
            REPO_DIR / candidate,
            Path.cwd() / candidate,
            BASE_DIR / candidate.name,
            REPO_DIR / candidate.name,
            Path.cwd() / candidate.name
        ]
        for c in candidates_to_try:
            if c.exists() and c.is_file():
                _add_file(c)
                break

    # 2. Files newly created or modified in monitored directories
    monitored_dirs = [OUTPUTS_DIR, GENERATED_DIR, BASE_DIR, REPO_DIR, Path.cwd()]
    for directory in monitored_dirs:
        if directory.exists():
            for f in directory.glob("*"):
                if f.is_file() and f.suffix.lower() in GENERATED_DOC_EXTENSIONS:
                    resolved_f = f.resolve()
                    if str(resolved_f) in seen_paths:
                        continue
                    mtime = f.stat().st_mtime
                    is_new = resolved_f not in before_snapshot
                    is_modified = mtime > (before_snapshot.get(resolved_f, 0.0) + 0.001)
                    is_recent = turn_start_time is not None and mtime >= (turn_start_time - 2.0)
                    if is_new or is_modified or is_recent:
                        _add_file(f)

    # 3. Fallback regex detection for file paths in response text
    patterns = [
        r'DOCUMENT_CREATED:\s*([^\s\n\r"\'<>]+)',
        r'(?:created|saved|generated)\s+(?:successfully\s+)?(?:at|to|as)?\s*:?\s*([^\n\r"\'<>]+\.(?:docx|xlsx|pdf|pptx|png|jpg|csv|txt))',
        r'(?:path|file|document):\s*([^\n\r"\'<>]+\.(?:docx|xlsx|pdf|pptx|png|jpg|csv|txt))',
        r'(?:^|[\s"\'`])([a-zA-Z0-9_\-\\\/\.]+\.(?:docx|xlsx|pdf|pptx|csv))\b'
    ]
    for pattern in patterns:
        for match in re.findall(pattern, response_text, re.IGNORECASE):
            clean = match.strip().strip("'\"`")
            if clean.upper().startswith("DOCUMENT_CREATED:"):
                clean = clean.split(":", 1)[1].strip()
            candidate = Path(clean)
            candidates_to_try = [
                candidate,
                OUTPUTS_DIR / candidate.name,
                GENERATED_DIR / candidate.name,
                BASE_DIR / candidate,
                REPO_DIR / candidate,
                Path.cwd() / candidate,
                BASE_DIR / candidate.name,
                REPO_DIR / candidate.name,
                Path.cwd() / candidate.name
            ]
            for c in candidates_to_try:
                if c.exists() and c.is_file():
                    _add_file(c)
                    break

    return detected_files


def _is_doc_generation_query(query: str) -> bool:
    """Checks if the user prompt is asking to generate or create a document or file."""
    q = query.lower()
    gen_words = ["create", "generate", "make", "build", "write", "save", "produce", "export", "download", "prepare"]
    doc_words = ["docx", "doc", "word", "excel", "xlsx", "spreadsheet", "sheet", "csv", "pdf", "powerpoint", "pptx", "presentation", "slides", "document", "file", "report"]
    has_gen = any(w in q for w in gen_words)
    has_doc = any(w in q for w in doc_words)
    return has_gen and has_doc


def _auto_generate_fallback_document(query: str, response_text: str) -> Optional[Path]:
    """
    Fallback safety net: if the user requested a document but the model did not
    invoke a tool and only outputted text, synthesize the document into OUTPUTS_DIR.
    """
    q_lower = query.lower()
    r_lower = response_text.lower()
    combined = q_lower + " " + r_lower

    # Detect extension
    ext = ".docx"
    if any(k in combined for k in [".xlsx", "excel", "spreadsheet"]):
        ext = ".xlsx"
    elif any(k in combined for k in [".pptx", "powerpoint", "presentation", "slides"]):
        ext = ".pptx"
    elif any(k in combined for k in [".pdf"]):
        ext = ".pdf"
    elif any(k in combined for k in [".docx", ".doc", "word", "document", "report"]):
        ext = ".docx"

    # Detect filename from query or response
    filename = None
    fn_match = re.search(r'([a-zA-Z0-9_\-]+\.(?:docx|xlsx|pptx|pdf))', query + " " + response_text, re.IGNORECASE)
    if fn_match:
        filename = fn_match.group(1).strip()
    else:
        title_match = re.search(r'titled\s+["\']?([a-zA-Z0-9_\-\s]+)["\']?', query + " " + response_text, re.IGNORECASE)
        if title_match:
            candidate_stem = title_match.group(1).strip().replace(" ", "_")
            if candidate_stem:
                filename = f"{candidate_stem}{ext}"
        if not filename:
            clean_name = re.sub(r'[^a-zA-Z0-9_]', '_', query[:25].strip()).strip('_')
            filename = f"{clean_name if clean_name else 'Document'}{ext}"

    if not filename.lower().endswith(ext):
        filename = f"{Path(filename).stem}{ext}"

    target_path = OUTPUTS_DIR / filename

    lines = [line.strip() for line in response_text.split("\n") if line.strip()]
    raw_title = lines[0] if lines else "Generated Document"
    clean_title = re.sub(r'^[#*_\-\s]+', '', raw_title)[:80]
    content = response_text

    try:
        if ext == ".docx":
            from docx import Document
            doc = Document()
            doc.add_heading(clean_title, level=1)
            for p in content.split("\n"):
                if p.strip():
                    doc.add_paragraph(p.strip())
            target_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(target_path)

        elif ext == ".xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Report"
            ws["A1"] = clean_title
            row_idx = 2
            for line in lines:
                cols = [c.strip() for c in line.split("|")]
                for col_idx, val in enumerate(cols, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=val)
                row_idx += 1
            target_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(target_path)

        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = clean_title
            tf = slide.placeholders[1].text_frame
            tf.clear()
            for line in lines:
                if line.strip():
                    p = tf.add_paragraph()
                    p.text = line.strip()
            target_path.parent.mkdir(parents=True, exist_ok=True)
            prs.save(target_path)

        elif ext == ".pdf":
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            target_path.parent.mkdir(parents=True, exist_ok=True)
            doc = SimpleDocTemplate(str(target_path), pagesize=A4)
            styles = getSampleStyleSheet()
            elements = [Paragraph(clean_title, styles["Title"]), Spacer(1, 20)]
            for line in lines:
                elements.append(Paragraph(line, styles["BodyText"]))
                elements.append(Spacer(1, 10))
            doc.build(elements)

        if target_path.exists() and target_path.is_file():
            return target_path.resolve()
    except Exception as e:
        pass

    return None


@router.post("/chat", response_model=ChatResponse)
def chat_endpoint(payload: ChatRequest):
    """
    Main agent endpoint for the Java backend.
    Processes the query through query analysis, model routing,
    agent tool execution, and verification.
    Returns:
    - AI response text
    - Required tools and tools actually executed
    - Full AgentState snapshot
    - Created files with filename, directory, full file_path, and download_url
    - Primary output_directory path
    """
    try:
        config = {"configurable": {"thread_id": payload.thread_id}}
        turn_start_time = datetime.now().timestamp()

        # Snapshot files before execution across potential output locations
        monitored_dirs = [OUTPUTS_DIR, GENERATED_DIR, BASE_DIR, REPO_DIR, Path.cwd()]
        before_files = _get_directory_files_snapshot(monitored_dirs)

        # Prepare query and user content if file_path is supplied
        query_text = payload.query
        file_path_str = payload.file_path.strip() if payload.file_path else ""

        # If user asks to create a document, instruct the model clearly to use document creation tools
        if _is_doc_generation_query(payload.query):
            doc_instruction = (
                f"\nIMPORTANT INSTRUCTION: You MUST use the appropriate available document creation tool "
                f"(create_docx, create_excel, create_pptx, or create_pdf) to generate this document locally on disk. "
                f"Do not simply reply with text."
            )
            query_text = f"{payload.query}{doc_instruction}"

        if file_path_str:
            file_ext = Path(file_path_str).suffix.lower()
            if file_ext in [".jpg", ".jpeg", ".png", ".webp"]:
                user_content = f"{payload.query}\nAttached image file: {file_path_str}"
            else:
                query_text = f"{payload.query} (Document file: {file_path_str})"
                user_content = (
                    f"User task: {payload.query}\n"
                    f"Attached local file: {file_path_str}\n"
                    f"Use the appropriate tool (such as read_text_file or analyze_dataset) to examine this file."
                )
        else:
            user_content = query_text

        # Invoke the LangGraph workflow
        result = graph.invoke({
            "query": query_text,
            "file_path": file_path_str,
            "verification_attempts": 0,
            "messages": [{"role": "user", "content": user_content}]
        }, config=config)

        raw_response = result.get("response", "No response generated.")

        # Extract tools executed from message history
        messages_history = result.get("messages", [])
        tools_used = _extract_tools_used(messages_history)

        # Extract created file paths recorded in AgentState
        state_created_files = list(result.get("created_files", []))

        # Also inspect messages history for any DOCUMENT_CREATED output
        for msg in messages_history:
            content = getattr(msg, "content", "")
            if isinstance(content, str) and "DOCUMENT_CREATED:" in content:
                for match in re.findall(r"DOCUMENT_CREATED:\s*([^\s\n\r\"']+)", content):
                    if match not in state_created_files:
                        state_created_files.append(match)

        # Resolve created files to metadata objects with full directory info
        created_files_info = _resolve_created_files(
            state_created_files,
            before_files,
            raw_response,
            turn_start_time=turn_start_time
        )

        # Extract analysis details
        analysis_raw = result.get("analysis")
        if hasattr(analysis_raw, "model_dump"):
            analysis_dict = analysis_raw.model_dump()
        elif isinstance(analysis_raw, dict):
            analysis_dict = analysis_raw
        else:
            analysis_dict = None

        # Safety Fallback: If no files were detected on disk but document generation was requested or claimed,
        # synthesize the document directly into OUTPUTS_DIR
        if not created_files_info:
            needs_doc = (
                _is_doc_generation_query(payload.query)
                or (analysis_dict and analysis_dict.get("task_type") == "document_generation")
                or (analysis_dict and analysis_dict.get("output_format") in ["docx", "xlsx", "pptx", "pdf"])
                or any(k in raw_response.lower() for k in [
                    "i have created", "i have generated", "saved as", "saved to",
                    "document has been created", "docx document", "excel file", "pdf document"
                ])
            )
            if needs_doc:
                fallback_file = _auto_generate_fallback_document(payload.query, raw_response)
                if fallback_file and fallback_file.exists():
                    f_dir = str(fallback_file.parent.resolve())
                    f_path = str(fallback_file.resolve())
                    created_files_info.append(
                        CreatedFileInfo(
                            filename=fallback_file.name,
                            directory=f_dir,
                            file_directory=f_dir,
                            file_path=f_path,
                            output_directory=f_dir,
                            output_file_directory=f_dir,
                            download_url=f"/api/download/{fallback_file.name}",
                            size_bytes=fallback_file.stat().st_size
                        )
                    )

        canonical_output_dir = str(OUTPUTS_DIR.resolve())

        # Collect resolved absolute file paths and their parent directories
        created_file_paths: List[str] = [cf.file_path for cf in created_files_info]
        created_file_dirs: List[str] = list(dict.fromkeys([cf.directory for cf in created_files_info if cf.directory]))

        # Fallback if state_created_files has paths not captured in created_files_info
        if not created_file_paths and state_created_files:
            for scf in state_created_files:
                clean_scf = str(scf).strip().strip("'\"")
                if clean_scf.upper().startswith("DOCUMENT_CREATED:"):
                    clean_scf = clean_scf.split(":", 1)[1].strip()
                cand = Path(clean_scf)
                if cand.is_file():
                    abs_path = str(cand.resolve())
                    abs_dir = str(cand.parent.resolve())
                    if abs_path not in created_file_paths:
                        created_file_paths.append(abs_path)
                    if abs_dir not in created_file_dirs:
                        created_file_dirs.append(abs_dir)

        # Determine the primary directory for the generated file(s)
        primary_file_dir = created_file_dirs[0] if created_file_dirs else (canonical_output_dir if created_file_paths else None)

        # Infer tools used from created files if not already in tools_used
        for cf in created_files_info:
            suffix = Path(cf.filename).suffix.lower()
            tool_mapping = {
                ".docx": "create_docx",
                ".xlsx": "create_excel",
                ".pptx": "create_pptx",
                ".pdf": "create_pdf"
            }
            mapped_tool = tool_mapping.get(suffix)
            if mapped_tool and mapped_tool not in tools_used:
                tools_used.append(mapped_tool)

        # Extract analysis details
        analysis_raw = result.get("analysis")
        if hasattr(analysis_raw, "model_dump"):
            analysis_dict = analysis_raw.model_dump()
        elif isinstance(analysis_raw, dict):
            analysis_dict = analysis_raw
        else:
            analysis_dict = None

        required_tools = None
        if analysis_dict:
            required_tools = bool(analysis_dict.get("needs_tools", False))

        # Build full AgentState snapshot for Java backend inspection
        agent_state_info = AgentStateInfo(
            model_role=result.get("model_role"),
            model_name=result.get("model_name"),
            analysis=analysis_dict,
            verified=result.get("verified"),
            verification_reason=result.get("verification_reason"),
            verification_attempts=result.get("verification_attempts"),
            tools_used=tools_used,
            created_files=created_file_paths,
            created_file_directories=created_file_dirs,
            directory=primary_file_dir,
            file_directory=primary_file_dir,
            output_directory=canonical_output_dir,
            output_file_directory=primary_file_dir
        )

        return ChatResponse(
            query=payload.query,
            response=raw_response,
            model_role=result.get("model_role"),
            model_name=result.get("model_name"),
            verified=result.get("verified"),
            verification_reason=result.get("verification_reason"),
            verification_attempts=result.get("verification_attempts"),
            tools_used=tools_used,
            required_tools=required_tools,
            analysis=analysis_dict,
            created_files=created_files_info,
            output_directory=canonical_output_dir,
            file_directory=primary_file_dir,
            output_file_directory=primary_file_dir,
            created_file_directory=primary_file_dir,
            created_file_paths=created_file_paths,
            created_file_directories=created_file_dirs,
            thread_id=payload.thread_id,
            agent_state=agent_state_info
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Agent workflow error: {str(e)}")
